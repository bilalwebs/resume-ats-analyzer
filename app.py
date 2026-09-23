import io
import json
import os
import re
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader


# ============================================================
# CONFIGURATION
# ============================================================

MODEL = "gemini-2.5-flash"
MAX_FILE_SIZE_MB = 10


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)


# ============================================================
# HEADER / BRANDING
# ============================================================

st.title("📄 Resume ATS Analyzer")

st.caption(
    "AI-powered ATS readiness analysis using Gemini 2.5 Flash"
)

st.markdown(
    "**Created by Muhammad Bilal Hussain**  \n"
    "Full Stack Developer & AI Engineer"
)

st.info(
    "ATS scores are estimates, not scores from a specific employer's ATS. "
    "Different ATS platforms and job descriptions use different rules."
)


# ============================================================
# API KEY
# ============================================================

def get_api_key() -> str | None:
    """
    Get Gemini API key from Streamlit Secrets first,
    then environment variables.

    The API key is never requested from the end user.
    """

    try:
        key = st.secrets.get("GEMINI_API_KEY")

        if key:
            return str(key).strip()

    except Exception:
        pass

    key = os.getenv("GEMINI_API_KEY")

    return key.strip() if key else None


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_text(uploaded_file: Any) -> str:
    """
    Extract text from PDF, DOCX, or TXT files.
    """

    extension = uploaded_file.name.lower().rsplit(".", 1)[-1]

    data = uploaded_file.getvalue()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == "pdf":

        reader = PdfReader(
            io.BytesIO(data)
        )

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == "docx":

        document = Document(
            io.BytesIO(data)
        )

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:

            for row in table.rows:

                parts.append(
                    " | ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    return data.decode(
        "utf-8",
        errors="replace"
    )


# ============================================================
# JSON PARSER
# ============================================================

def parse_json(text: str) -> dict[str, Any]:
    """
    Safely parse Gemini's JSON response.
    """

    if not text:
        raise ValueError(
            "Gemini returned an empty response."
        )

    cleaned = text.strip()

    # Remove Markdown JSON fences if Gemini returns them.
    cleaned = re.sub(
        r"^```(?:json)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )

    cleaned = re.sub(
        r"\s*```$",
        "",
        cleaned,
    )

    match = re.search(
        r"\{.*\}",
        cleaned,
        flags=re.DOTALL,
    )

    if not match:

        raise ValueError(
            "Gemini did not return valid JSON."
        )

    result = json.loads(
        match.group(0)
    )

    if not isinstance(result, dict):

        raise ValueError(
            "Gemini returned an unexpected response format."
        )

    return result


# ============================================================
# GEMINI ATS ANALYSIS
# ============================================================

def analyze_resume(
    resume_text: str,
    job_description: str,
) -> dict[str, Any]:

    api_key = get_api_key()

    if not api_key:

        raise RuntimeError(
            "GEMINI_API_KEY is not configured. "
            "Please add it to Streamlit Secrets."
        )

    client = genai.Client(
        api_key=api_key
    )

    target_job = (
        job_description.strip()
        if job_description.strip()
        else (
            "Not provided. "
            "Evaluate general ATS readiness."
        )
    )

    prompt = f"""
You are a professional resume and ATS optimization analyst.

Analyze the resume against the job description when one is provided.

IMPORTANT RULES:

1. Do not invent experience.
2. Do not invent education.
3. Do not invent skills.
4. Do not invent employers.
5. Do not invent projects.
6. Do not invent metrics.
7. Do not invent achievements.
8. Recommendations must remain truthful to the resume.
9. The ATS score is an estimate, not the score of a specific ATS product.
10. Do not recommend adding a keyword if the resume provides no truthful basis for it.
11. Clearly distinguish missing keywords from skills the candidate actually possesses.
12. Focus on practical, actionable resume improvements.

Score the resume from 0 to 100 using these dimensions:

- ATS-readable structure
- Formatting
- Standard resume sections
- Keyword alignment
- Skills relevance
- Experience relevance
- Achievement clarity
- Measurable impact
- Date consistency
- Job-title consistency
- Overall content clarity

If a job description is not provided:

Evaluate general ATS readiness and clearly explain that
keyword matching is limited.

Return ONLY valid JSON.

Use exactly this structure:

{{
  "ats_score": 0,
  "score_label": "Strong",
  "summary": "...",

  "keyword_match": {{
    "score": 0,
    "matched_keywords": ["..."],
    "missing_keywords": ["..."],
    "notes": "..."
  }},

  "sections": [
    {{
      "section": "Experience",
      "status": "Good",
      "feedback": "..."
    }}
  ],

  "formatting": {{
    "score": 0,
    "strengths": ["..."],
    "issues": ["..."]
  }},

  "content": {{
    "score": 0,
    "strengths": ["..."],
    "issues": ["..."]
  }},

  "improvements": [
    {{
      "priority": "High",
      "area": "Professional Summary",
      "problem": "...",
      "recommendation": "..."
    }}
  ],

  "quick_wins": [
    "..."
  ],

  "ats_checklist": [
    {{
      "item": "Standard section headings",
      "status": "Pass",
      "explanation": "..."
    }}
  ]
}}

Additional requirements:

- Keep all scores between 0 and 100.
- Keep improvements focused on approximately 5 to 10 genuinely useful recommendations.
- Use concise but specific feedback.
- Never fabricate candidate information.
- Do not claim that this is an official ATS score.
- Analyze only the resume content provided.

JOB DESCRIPTION:

{target_job}


RESUME:

{resume_text}
"""

    response = client.models.generate_content(
        model=MODEL,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    return parse_json(
        response.text or ""
    )


# ============================================================
# SCORE VALIDATION
# ============================================================

def clamp_score(value: Any) -> int:

    try:

        return max(
            0,
            min(
                100,
                int(value),
            ),
        )

    except (
        TypeError,
        ValueError,
    ):

        return 0


# ============================================================
# RESULTS UI
# ============================================================

def render_results(
    result: dict[str, Any],
) -> None:

    ats_score = clamp_score(
        result.get(
            "ats_score",
            0,
        )
    )

    score_label = (
        str(
            result.get(
                "score_label",
                "Needs work",
            )
        ).strip()
        or "Needs work"
    )

    summary = (
        str(
            result.get(
                "summary",
                "No summary provided.",
            )
        ).strip()
        or "No summary provided."
    )

    keyword_match = (
        result.get(
            "keyword_match"
        )
        if isinstance(
            result.get("keyword_match"),
            dict,
        )
        else {}
    )

    formatting = (
        result.get(
            "formatting"
        )
        if isinstance(
            result.get("formatting"),
            dict,
        )
        else {}
    )

    content = (
        result.get(
            "content"
        )
        if isinstance(
            result.get("content"),
            dict,
        )
        else {}
    )

    sections = (
        result.get(
            "sections"
        )
        if isinstance(
            result.get("sections"),
            list,
        )
        else []
    )

    improvements = (
        result.get(
            "improvements"
        )
        if isinstance(
            result.get("improvements"),
            list,
        )
        else []
    )

    quick_wins = (
        result.get(
            "quick_wins"
        )
        if isinstance(
            result.get("quick_wins"),
            list,
        )
        else []
    )

    checklist = (
        result.get(
            "ats_checklist"
        )
        if isinstance(
            result.get("ats_checklist"),
            list,
        )
        else []
    )

    # ========================================================
    # ATS OVERVIEW
    # ========================================================

    st.subheader("📊 ATS Overview")

    col1, col2, col3 = st.columns(
        [1.5, 1.3, 1.3]
    )

    with col1:

        st.metric(
            "ATS Score",
            f"{ats_score}/100",
        )

    with col2:

        st.metric(
            "Assessment",
            score_label,
        )

    with col3:

        st.metric(
            "Keyword Match",
            f"{clamp_score(keyword_match.get('score', 0))}/100",
        )

    st.progress(
        ats_score / 100
    )

    # ========================================================
    # SUMMARY
    # ========================================================

    st.subheader(
        "🧾 Overall Assessment"
    )

    st.write(summary)

    # ========================================================
    # KEYWORD ANALYSIS
    # ========================================================

    st.subheader(
        "🔑 Keyword Analysis"
    )

    matched = keyword_match.get(
        "matched_keywords",
        [],
    )

    missing = keyword_match.get(
        "missing_keywords",
        [],
    )

    notes = keyword_match.get(
        "notes",
        "",
    )

    col4, col5 = st.columns(2)

    with col4:

        st.markdown(
            "**Matched Keywords**"
        )

        if matched:

            st.write(
                ", ".join(
                    str(item)
                    for item in matched
                )
            )

        else:

            st.write(
                "No strongly matched keywords identified."
            )

    with col5:

        st.markdown(
            "**Missing / Useful Keywords**"
        )

        if missing:

            st.write(
                ", ".join(
                    str(item)
                    for item in missing
                )
            )

        else:

            st.write(
                "No obvious keyword gaps identified."
            )

    if notes:

        st.caption(notes)

    # ========================================================
    # SECTION REVIEW
    # ========================================================

    st.subheader(
        "📌 Section Review"
    )

    for item in sections:

        section_name = (
            str(
                item.get(
                    "section",
                    "Section",
                )
            ).strip()
            or "Section"
        )

        status = (
            str(
                item.get(
                    "status",
                    "Needs review",
                )
            ).strip()
            or "Needs review"
        )

        feedback = (
            str(
                item.get(
                    "feedback",
                    "No feedback provided.",
                )
            ).strip()
            or "No feedback provided."
        )

        with st.expander(
            f"{section_name} — {status}"
        ):

            st.write(
                feedback
            )

    # ========================================================
    # FORMATTING + CONTENT
    # ========================================================

    st.subheader(
        "🛠️ Formatting & Content"
    )

    col6, col7 = st.columns(2)

    with col6:

        st.markdown(
            f"**Formatting Score: "
            f"{clamp_score(formatting.get('score', 0))}/100**"
        )

        strengths = formatting.get(
            "strengths",
            [],
        )

        issues = formatting.get(
            "issues",
            [],
        )

        if strengths:

            st.write("**Strengths:**")

            for point in strengths:

                st.write(
                    f"✅ {point}"
                )

        if issues:

            st.write("**Issues:**")

            for point in issues:

                st.write(
                    f"⚠️ {point}"
                )

    with col7:

        st.markdown(
            f"**Content Score: "
            f"{clamp_score(content.get('score', 0))}/100**"
        )

        strengths = content.get(
            "strengths",
            [],
        )

        issues = content.get(
            "issues",
            [],
        )

        if strengths:

            st.write("**Strengths:**")

            for point in strengths:

                st.write(
                    f"✅ {point}"
                )

        if issues:

            st.write("**Issues:**")

            for point in issues:

                st.write(
                    f"⚠️ {point}"
                )

    # ========================================================
    # PRIORITY IMPROVEMENTS
    # ========================================================

    st.subheader(
        "🚀 Recommended Improvements"
    )

    if improvements:

        for index, item in enumerate(
            improvements,
            start=1,
        ):

            priority = (
                str(
                    item.get(
                        "priority",
                        "Medium",
                    )
                ).strip()
                or "Medium"
            )

            area = (
                str(
                    item.get(
                        "area",
                        "General",
                    )
                ).strip()
                or "General"
            )

            problem = (
                str(
                    item.get(
                        "problem",
                        "No problem identified.",
                    )
                ).strip()
                or "No problem identified."
            )

            recommendation = (
                str(
                    item.get(
                        "recommendation",
                        "No recommendation provided.",
                    )
                ).strip()
                or "No recommendation provided."
            )

            with st.expander(
                f"{index}. {area} — {priority}"
            ):

                st.markdown(
                    f"**Problem:** {problem}"
                )

                st.markdown(
                    f"**Recommendation:** {recommendation}"
                )

    else:

        st.write(
            "No improvement items were returned."
        )

    # ========================================================
    # QUICK WINS
    # ========================================================

    st.subheader(
        "⚡ Quick Wins"
    )

    if quick_wins:

        for item in quick_wins:

            st.write(
                f"• {item}"
            )

    else:

        st.write(
            "No quick wins were identified."
        )

    # ========================================================
    # ATS CHECKLIST
    # ========================================================

    st.subheader(
        "✅ ATS Checklist"
    )

    for item in checklist:

        label = (
            str(
                item.get(
                    "item",
                    "Checklist item",
                )
            ).strip()
            or "Checklist item"
        )

        status = (
            str(
                item.get(
                    "status",
                    "Unknown",
                )
            ).strip()
            or "Unknown"
        )

        explanation = (
            str(
                item.get(
                    "explanation",
                    "No explanation provided.",
                )
            ).strip()
            or "No explanation provided."
        )

        status_lower = status.lower()

        icon = (
            "✅"
            if status_lower in {
                "pass",
                "good",
                "yes",
            }
            else "⚠️"
        )

        st.write(
            f"{icon} **{label}** — "
            f"{status} — "
            f"{explanation}"
        )

    # ========================================================
    # RAW JSON
    # ========================================================

    with st.expander(
        "View Raw Analysis JSON"
    ):

        st.json(result)


# ============================================================
# FILE UPLOAD
# ============================================================

uploaded_file = st.file_uploader(
    "📎 Upload Your Resume",
    type=[
        "pdf",
        "docx",
        "txt",
    ],
    help=(
        "Supported formats: PDF, DOCX, TXT. "
        "Maximum size: 10 MB."
    ),
)


# ============================================================
# JOB DESCRIPTION
# ============================================================

job_description = st.text_area(
    "💼 Target Job Description",
    height=220,
    placeholder=(
        "Paste the target job description here "
        "for targeted keyword analysis..."
    ),
)


st.caption(
    "💡 Adding a job description provides more targeted "
    "keyword and relevance analysis."
)


# ============================================================
# ANALYZE BUTTON
# ============================================================

if st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
):

    # --------------------------------------------------------
    # File validation
    # --------------------------------------------------------

    if uploaded_file is None:

        st.warning(
            "Please upload a resume first."
        )

        st.stop()

    # --------------------------------------------------------
    # File size validation
    # --------------------------------------------------------

    if (
        uploaded_file.size
        > MAX_FILE_SIZE_MB * 1024 * 1024
    ):

        st.error(
            f"File is too large. "
            f"Please upload a file under "
            f"{MAX_FILE_SIZE_MB} MB."
        )

        st.stop()

    try:

        # ----------------------------------------------------
        # Extract resume
        # ----------------------------------------------------

        with st.spinner(
            "📄 Extracting resume text..."
        ):

            resume_text = extract_text(
                uploaded_file
            )

        if len(
            resume_text.strip()
        ) < 80:

            st.error(
                "Very little readable text was found. "
                "If this is a scanned PDF, upload a "
                "text-based PDF or DOCX instead."
            )

            st.stop()

        # ----------------------------------------------------
        # Gemini analysis
        # ----------------------------------------------------

        with st.spinner(
            "🤖 Running Gemini ATS analysis..."
        ):

            analysis = analyze_resume(
                resume_text,
                job_description,
            )

        # ----------------------------------------------------
        # Save result
        # ----------------------------------------------------

        st.session_state.analysis = analysis

        st.session_state.resume_name = (
            uploaded_file.name
        )

        st.success(
            "Resume analysis completed successfully."
        )

    except Exception as exc:

        st.error(
            f"Analysis failed: {exc}"
        )

        with st.expander(
            "Technical Error Details"
        ):

            st.exception(exc)

        st.stop()


# ============================================================
# DISPLAY SAVED RESULTS
# ============================================================

result = st.session_state.get(
    "analysis"
)

if result:

    st.divider()

    st.subheader(
        f"📋 Analysis — "
        f"{st.session_state.get('resume_name', 'Resume')}"
    )

    render_results(
        result
    )


# ============================================================
# PRIVACY NOTICE
# ============================================================

st.divider()

st.caption(
    "🔐 Privacy: Resume text is sent to Google's "
    "Gemini API for analysis. Do not upload confidential "
    "documents unless you are comfortable with that processing."
)


# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """
    <div style="text-align: center; padding: 20px 0 10px 0;">
        <p style="margin-bottom: 4px;">
            <strong>Resume ATS Analyzer</strong>
        </p>
        <p style="margin-bottom: 4px;">
            Built by <strong>Muhammad Bilal Hussain</strong>
        </p>
        <p style="font-size: 13px; color: #888;">
            Full Stack Developer & AI Engineer
            • Powered by Gemini 2.5 Flash
        </p>
    </div>
    """,
    unsafe_allow_html=True,
)
